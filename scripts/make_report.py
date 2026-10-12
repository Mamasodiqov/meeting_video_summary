# -*- coding: utf-8 -*-
"""Build the formal DOCX report of the meeting Q&A answers."""
import os
import sys

sys.path.insert(0, os.path.dirname(__file__))
from report_content import META, PARTICIPANTS, QUESTIONS, METHODOLOGY  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402

OUT = os.path.join(os.path.dirname(__file__), "..", "reports",
                   "meeting_qa_hisobot.docx")


def set_font(style_or_run, size=12, bold=False, italic=False):
    f = style_or_run.font
    f.name = "Times New Roman"
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = RGBColor(0, 0, 0)
    rpr = style_or_run.element.get_or_add_rPr() if hasattr(
        style_or_run.element, "get_or_add_rPr") else None
    if rpr is not None:
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), "Times New Roman")


def add_para(doc, text="", size=12, bold=False, italic=False,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, indent=None):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if indent is not None:
        p.paragraph_format.left_indent = Pt(indent)
    if text:
        r = p.add_run(text)
        set_font(r, size=size, bold=bold, italic=italic)
    return p


def add_quote(doc, ts, speaker, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Pt(36)
    p.paragraph_format.space_after = Pt(8)
    r1 = p.add_run(f"[{ts}] {speaker}: ")
    set_font(r1, size=12, bold=True)
    r2 = p.add_run(text)
    set_font(r2, size=12)


def main():
    doc = Document()

    normal = doc.styles["Normal"]
    set_font(normal, size=12)

    # Title
    add_para(doc, "O‘QIM SAVOLLARIGA BERILGAN JAVOBLAR BO‘YICHA",
             size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    add_para(doc, "HISOBOT", size=14, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_para(doc, "«Cloudflare» platformasini joriy etish bo‘yicha",
             size=12, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER,
             space_after=12)

    # Meta block
    meta_lines = [
        ("O‘qim sanasi:", META["meeting_date"]),
        ("O‘qim davomiyligi:", META["duration"]),
        ("Manba:", META["source"]),
        ("Hisobot tuzilgan sana:", META["made_date"]),
    ]
    for k, v in meta_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r1 = p.add_run(f"{k} ")
        set_font(r1, bold=True)
        r2 = p.add_run(v)
        set_font(r2)
    add_para(doc, "", space_after=6)

    # 1. Qisqacha
    h = doc.add_heading("", level=1)
    r = h.add_run("1. O‘qim haqida qisqacha")
    set_font(r, size=13, bold=True)
    add_para(doc, PARTICIPANTS)

    # 2. Savollar va javoblar
    h = doc.add_heading("", level=1)
    r = h.add_run("2. Savollar va o‘qimda berilgan javoblar")
    set_font(r, size=13, bold=True)
    add_para(doc, "Quyida markaz tomonidan yuborilgan 5 ta savol va o‘qim "
                  "davomida berilgan javoblar keltirilgan. Javoblar "
                  "transkripsiyadan o‘zgarishsiz, asl tilida iqtibos "
                  "qilib keltiriladi; har bir iqtibos yonida o‘qim "
                  "vaqtidagi mos bo‘lim (ss:ss) ko‘rsatilgan.")

    for i, item in enumerate(QUESTIONS, 1):
        h = doc.add_heading("", level=2)
        r = h.add_run(f"2.{i} Savol: {item['q']}")
        set_font(r, size=12, bold=True)

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        r1 = p.add_run("Holat: ")
        set_font(r1, bold=True)
        r2 = p.add_run(item["status"] + ".")
        set_font(r2)

        if item.get("note"):
            add_para(doc, item["note"], italic=True, space_after=8)

        add_para(doc, "O‘qimda aytilgan javob (iqtibos):", bold=True,
                 space_after=4)
        for ts, spk, txt in item["quotes"]:
            add_quote(doc, ts, spk, txt)
        if item.get("note2"):
            add_para(doc, item["note2"], italic=True, size=11, space_after=10)

    # 3. Metodika
    h = doc.add_heading("", level=1)
    r = h.add_run("3. Metodika va sifat haqida eslatma")
    set_font(r, size=13, bold=True)
    add_para(doc, METHODOLOGY)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)
    print("saved:", os.path.abspath(OUT))


if __name__ == "__main__":
    main()
